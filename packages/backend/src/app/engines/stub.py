"""stub 引擎（替身）：不调用真实 AI，仅生成占位数据，用于打通整条链路。

实现 PptEngine 协议（与 packages/engine/core/protocol.py 一致），
因此后端 service 层只认协议，stub 与真实 PptMasterEngine 可无缝替换（"引擎可插拔"）。
"""
from __future__ import annotations

import json
import sys
import threading
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

# 确保 engine 包可被 import（直接 import 本文件时也安全）
_PACKAGES_DIR = Path(__file__).resolve().parents[4]
if str(_PACKAGES_DIR) not in sys.path:
    sys.path.insert(0, str(_PACKAGES_DIR))

from engine.core.config import get_jobs_root
from engine.core.models import Job, JobId, JobStatus
from engine.core.protocol import PptEngine


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


class StubPptMasterEngine(PptEngine):
    """ppt_master 引擎的替身实现（实现 PptEngine 协议）。"""

    def __init__(self, jobs_root=None, agent: str = "claude", simulate_seconds: int = 3):
        self.jobs_root: Path = (jobs_root or get_jobs_root()).resolve()
        self.jobs_root.mkdir(parents=True, exist_ok=True)
        self.agent = agent
        self.simulate_seconds = simulate_seconds

    # ── 内部工具 ──
    def _job_dir(self, job_id) -> Path:
        return self.jobs_root / str(job_id)

    def _read_job(self, job_id) -> dict:
        p = self._job_dir(job_id) / "job.json"
        if not p.exists():
            return {}
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            return {}

    def _write_job(self, job_id, **fields) -> dict:
        p = self._job_dir(job_id) / "job.json"
        data = self._read_job(job_id)
        data.update(fields)
        p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        return data

    # ── 草稿/产物构造（保留原 stub 能力，作为内部 helper）──
    @staticmethod
    def _outline_sections(content_style: str) -> list[str]:
        return {
            "product_story": ["## 用户痛点", "## 产品方案", "## 使用场景", "## 价值收益"],
            "briefing": ["## 背景", "## 现状", "## 方案", "## 计划"],
            "event_promo": ["## 活动亮点", "## 时间地点", "## 参与收益", "## 行动号召"],
            "investment_pitch": ["## 市场机会", "## 商业模式", "## 合作价值", "## 下一步"],
        }.get(content_style, ["## 目标受众", "## 核心卖点", "## 价值收益", "## 行动号召"])

    @classmethod
    def generate_draft(cls, title: str, pages: int, content: str | None, engine_params: dict | None = None) -> str:
        style = (engine_params or {}).get("style", "swiss-minimal")
        content_style = (engine_params or {}).get("content_style", "marketing_promo")
        body = content or "（请在确认前补充正文内容，或提供 source 素材）"
        lines = [
            f"# {title}",
            "",
            f"> 风格：{style} ｜ 内容风格：{content_style} ｜ 目标页数：{pages}",
            "",
            "## 摘要",
            body,
            "",
        ]
        outline = cls._outline_sections(content_style)
        for i in range(1, max(int(pages), 1) + 1):
            section = outline[(i - 1) % len(outline)]
            lines += [f"## 第 {i} 页", "", section, "", f"- 要点 {i}.1", f"- 要点 {i}.2", ""]
        lines += ["---", "_（这是 stub 生成的占位草稿，接入真实引擎后会被替换）_", ""]
        return "\n".join(lines)

    @staticmethod
    def generate_ppt(job_dir: Path, title: str, content: str | None) -> Path:
        job_dir.mkdir(parents=True, exist_ok=True)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "内容（stub 占位）"
        body = slide.placeholders[1]
        body.text = (content or "stub 占位内容")[:500]
        out = job_dir / "result.pptx"
        prs.save(out)
        return out

    # ── PptEngine 协议 ──
    def submit(self, content, title=None, pages=None, engine_params=None, source=None, job_id=None) -> JobId:
        job_id = job_id or uuid.uuid4().hex
        job_dir = self._job_dir(job_id)
        job_dir.mkdir(parents=True, exist_ok=True)
        text = content or (source.get("value") if source else "")
        
        # 使用默认值处理可选的 title 和 pages
        effective_title = title or "Untitled"
        effective_pages = pages or 10
        
        draft = self.generate_draft(effective_title, effective_pages, text, engine_params)
        (job_dir / "draft.md").write_text(draft, encoding="utf-8")
        self._write_job(
            job_id,
            status="draft",
            title=effective_title,
            pages=int(effective_pages),
            engine_params=engine_params or {},
            created_at=_now(), updated_at=_now(), pid=None, result_path=None, error=None,
        )
        return JobId(job_id)

    def confirm(self, job_id, content, engine_params) -> None:
        job_dir = self._job_dir(job_id)
        job_dir.mkdir(parents=True, exist_ok=True)
        (job_dir / "draft.md").write_text(content or "", encoding="utf-8")
        self._write_job(job_id, status="running", pid=None, started_at=time.time(), updated_at=_now(), error=None)
        # 用后台线程模拟"长任务生成"，结束后把状态写回 job.json
        threading.Thread(target=self._run_generation, args=(job_id, content), daemon=True).start()

    def _run_generation(self, job_id, content) -> None:
        time.sleep(self.simulate_seconds)
        title = self._read_job(job_id).get("title", "")
        result = self.generate_ppt(self._job_dir(job_id), title, content)
        self._write_job(job_id, status="completed", result_path=str(result), updated_at=_now())

    def status(self, job_id) -> JobStatus:
        return JobStatus(self._read_job(job_id).get("status", "draft"))

    def result(self, job_id) -> str:
        rp = self._read_job(job_id).get("result_path")
        if rp and Path(rp).exists():
            return rp
        raise FileNotFoundError(f"任务 {job_id} 尚未生成产物")

    def resume(self, job_id) -> Job:
        data = self._read_job(job_id)
        if data.get("status") == "completed":
            return self._to_job(job_id, data)
        draft_path = self._job_dir(job_id) / "draft.md"
        content = draft_path.read_text(encoding="utf-8") if draft_path.exists() else ""
        self.confirm(job_id, content, data.get("engine_params") or {})
        return self._to_job(job_id, self._read_job(job_id))

    @staticmethod
    def _to_job(job_id, data) -> Job:
        return Job(
            job_id=JobId(str(job_id)),
            status=JobStatus(data.get("status", "draft")),
            result_path=data.get("result_path"),
            draft_path=str(data.get("draft_path")) if data.get("draft_path") else None,
            error=data.get("error"),
            created_at=data.get("created_at"),
            updated_at=data.get("updated_at"),
            content_style=data.get("content_style"),
        )

    def get_progress(self, job_id) -> dict | None:
        """获取任务进度信息，仅在 running 状态时返回。"""
        data = self._read_job(job_id)
        if data.get("status") != "running":
            return None
        
        started_at = data.get("started_at", 0)
        elapsed = int(time.time() - started_at) if started_at else 0
        
        return {
            "stage": "generating",
            "stage_label": "正在生成 PPT（stub 模拟）",
            "pages_generated": 2,
            "elapsed_seconds": elapsed,
            "current_page": "stub_page.pptx",
            "last_log": "stub 引擎模拟生成中...",
        }
